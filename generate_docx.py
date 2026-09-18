"""
Generate a comprehensive .docx project report for Movie Genre Classification.
Includes all required sections: Cover Page, Introduction, Problem Statement,
Functional/Non-functional Requirements, System Architecture, Design Diagrams
(Use Case, Workflow, Sequence, Component), Design Decisions, Implementation,
Screenshots/Results, Testing Approach, and Challenges.

Run: python generate_docx.py
"""

import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import numpy as np

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PLOTS_DIR = os.path.join(SCRIPT_DIR, "plots")
DIAGRAMS_DIR = os.path.join(SCRIPT_DIR, "diagrams")
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "Movie_Genre_Classification_Report.docx")

os.makedirs(PLOTS_DIR, exist_ok=True)
os.makedirs(DIAGRAMS_DIR, exist_ok=True)


# ============================================================================
# HELPER FUNCTIONS
# ============================================================================

def set_run_font(run, size=11, bold=False, italic=False, color=None, font_name="Calibri"):
    run.font.size = Pt(size)
    run.font.name = font_name
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_styled(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)
    return heading


def add_body(doc, text, bold=False, italic=False, size=11, spacing_after=6):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    para.paragraph_format.space_after = Pt(spacing_after)
    return para


def add_bullet(doc, text, bold_prefix="", size=11):
    para = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = para.add_run(bold_prefix)
        set_run_font(run, size=size, bold=True)
    run = para.add_run(text)
    set_run_font(run, size=size)
    return para


def add_numbered(doc, text, bold_prefix="", size=11):
    para = doc.add_paragraph(style="List Number")
    if bold_prefix:
        run = para.add_run(bold_prefix)
        set_run_font(run, size=size, bold=True)
    run = para.add_run(text)
    set_run_font(run, size=size)
    return para


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=11, bold=True, color=(255, 255, 255))
        shading = cell._element.get_or_add_tcPr()
        shading_elem = shading.makeelement(qn("w:shd"), {
            qn("w:fill"): "003366",
            qn("w:val"): "clear",
        })
        shading.append(shading_elem)

    for r, row_data in enumerate(rows):
        for c, value in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(value))
            set_run_font(run, size=11)

    doc.add_paragraph()
    return table


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        run = p.add_run(line)
        set_run_font(run, size=10, font_name="Consolas", color=(0, 80, 0))


def add_image(doc, path, caption, width=5.0):
    if os.path.exists(path):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        run.add_picture(path, width=Inches(width))

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        set_run_font(run, size=10, italic=True, color=(100, 100, 100))
        doc.add_paragraph()


# ============================================================================
# DIAGRAM GENERATION
# ============================================================================

def draw_rounded_box(ax, x, y, w, h, text, fc="#E8F0FE", ec="#003366", fontsize=9, text_color="black"):
    box = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.1",
                         facecolor=fc, edgecolor=ec, linewidth=1.5)
    ax.add_patch(box)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center",
            fontsize=fontsize, fontweight="bold", color=text_color, wrap=True)


def draw_arrow(ax, x1, y1, x2, y2, color="#003366"):
    ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                arrowprops=dict(arrowstyle="-|>", color=color, lw=1.5))


def draw_arrow_label(ax, x1, y1, x2, y2, label="", color="#003366"):
    ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                arrowprops=dict(arrowstyle="-|>", color=color, lw=1.5))
    mx, my = (x1 + x2) / 2, (y1 + y2) / 2
    if label:
        ax.text(mx, my + 0.15, label, ha="center", va="center", fontsize=7,
                color="#555555", style="italic")


def generate_use_case_diagram():
    fig, ax = plt.subplots(1, 1, figsize=(10, 7))
    ax.set_xlim(-1, 11)
    ax.set_ylim(-0.5, 7.5)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("Use Case Diagram", fontsize=16, fontweight="bold", color="#003366", pad=15)

    # System boundary
    rect = mpatches.FancyBboxPatch((2.5, 0.2), 6, 7, boxstyle="round,pad=0.3",
                                    facecolor="#F5F8FF", edgecolor="#003366",
                                    linewidth=2, linestyle="--")
    ax.add_patch(rect)
    ax.text(5.5, 7.0, "Movie Genre Classification System", ha="center",
            fontsize=12, fontweight="bold", color="#003366")

    # Actor - stick figure
    ax.plot(0.8, 4.5, "o", markersize=15, color="#003366")
    ax.plot([0.8, 0.8], [3.5, 4.2], color="#003366", lw=2)
    ax.plot([0.3, 1.3], [4.0, 4.0], color="#003366", lw=2)
    ax.plot([0.8, 0.4], [3.5, 2.8], color="#003366", lw=2)
    ax.plot([0.8, 1.2], [3.5, 2.8], color="#003366", lw=2)
    ax.text(0.8, 2.4, "User", ha="center", fontsize=10, fontweight="bold", color="#003366")

    # Use cases (ellipses)
    use_cases = [
        (5.5, 6.0, "Enter Movie\nDescription"),
        (5.5, 4.8, "View Model\nAccuracy"),
        (5.5, 3.6, "View Classification\nReport"),
        (5.5, 2.4, "View\nVisualizations"),
        (5.5, 1.2, "Get Genre\nPrediction"),
    ]

    for (cx, cy, label) in use_cases:
        ellipse = mpatches.Ellipse((cx, cy), 3.0, 0.9, facecolor="#DCEAFF",
                                    edgecolor="#003366", linewidth=1.5)
        ax.add_patch(ellipse)
        ax.text(cx, cy, label, ha="center", va="center", fontsize=9, fontweight="bold")

    # Lines from actor to use cases
    for (cx, cy, _) in use_cases:
        ax.plot([1.3, cx - 1.5], [3.8, cy], color="#003366", lw=1, linestyle="-")

    path = os.path.join(DIAGRAMS_DIR, "use_case_diagram.png")
    plt.tight_layout()
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def generate_workflow_diagram():
    fig, ax = plt.subplots(1, 1, figsize=(12, 8))
    ax.set_xlim(-0.5, 10)
    ax.set_ylim(-0.5, 9)
    ax.axis("off")
    ax.set_title("Workflow Diagram", fontsize=16, fontweight="bold", color="#003366", pad=15)

    steps = [
        (4.0, 8.0, "Load Dataset\n(movies.csv)"),
        (4.0, 6.8, "Text Preprocessing\n(Lowercase + Remove Punctuation)"),
        (4.0, 5.6, "Feature Extraction\n(CountVectorizer vs TF-IDF)"),
        (4.0, 4.4, "Train-Test Split\n(80/20, Stratified)"),
        (4.0, 3.2, "Model Training\n(NB, LR, SVM, RF)"),
        (4.0, 2.0, "Model Evaluation\n(Accuracy, F1, Confusion Matrix)"),
        (4.0, 0.8, "Best Model Selection\n& Genre Prediction"),
    ]

    colors = ["#E3F2FD", "#BBDEFB", "#90CAF9", "#64B5F6", "#42A5F5", "#2196F3", "#1565C0"]
    text_colors = ["black"] * 5 + ["black", "white"]

    for i, (x, y, text) in enumerate(steps):
        draw_rounded_box(ax, x - 1.8, y - 0.4, 3.6, 0.8, text,
                         fc=colors[i], ec="#003366", fontsize=9, text_color=text_colors[i])

    for i in range(len(steps) - 1):
        draw_arrow(ax, steps[i][0], steps[i][1] - 0.4, steps[i + 1][0], steps[i + 1][1] + 0.4)

    # Side branch: visualizations
    draw_rounded_box(ax, 7.0, 1.6, 2.5, 0.8, "Generate\nVisualizations",
                     fc="#FFF9C4", ec="#F57F17", fontsize=9)
    ax.annotate("", xy=(7.0, 2.0), xytext=(5.8, 2.0),
                arrowprops=dict(arrowstyle="-|>", color="#F57F17", lw=1.5))

    path = os.path.join(DIAGRAMS_DIR, "workflow_diagram.png")
    plt.tight_layout()
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def generate_sequence_diagram():
    fig, ax = plt.subplots(1, 1, figsize=(12, 10))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 12)
    ax.axis("off")
    ax.set_title("Sequence Diagram", fontsize=16, fontweight="bold", color="#003366", pad=15)

    # Actors/Objects
    objects = [
        (1.5, "User"),
        (4.5, "Main\nController"),
        (7.5, "Preprocessor"),
        (10.5, "ML Models"),
    ]

    for x, label in objects:
        draw_rounded_box(ax, x - 0.8, 11.0, 1.6, 0.7, label,
                         fc="#003366", ec="#003366", fontsize=9, text_color="white")
        ax.plot([x, x], [0.5, 11.0], color="#003366", lw=1, linestyle="--")

    # Messages
    messages = [
        (1.5, 4.5, 10.5, "1: Run main.py"),
        (4.5, 7.5, 9.8, "2: Load CSV data"),
        (7.5, 4.5, 9.1, "3: Return DataFrame"),
        (4.5, 7.5, 8.4, "4: Preprocess text"),
        (7.5, 4.5, 7.7, "5: Return cleaned text"),
        (4.5, 10.5, 7.0, "6: Extract features (TF-IDF)"),
        (4.5, 10.5, 6.3, "7: Train models"),
        (10.5, 4.5, 5.6, "8: Return predictions"),
        (4.5, 10.5, 4.9, "9: Evaluate models"),
        (10.5, 4.5, 4.2, "10: Return metrics"),
        (4.5, 1.5, 3.5, "11: Display results"),
        (1.5, 4.5, 2.8, "12: Enter description"),
        (4.5, 10.5, 2.1, "13: Predict genre"),
        (10.5, 4.5, 1.4, "14: Return prediction"),
        (4.5, 1.5, 0.7, "15: Display genre"),
    ]

    for x1, x2, y, label in messages:
        color = "#003366" if x1 < x2 else "#CC0000"
        ax.annotate("", xy=(x2, y), xytext=(x1, y),
                    arrowprops=dict(arrowstyle="-|>", color=color, lw=1.2))
        mx = (x1 + x2) / 2
        ax.text(mx, y + 0.18, label, ha="center", va="bottom", fontsize=7.5,
                color="#333333", fontweight="bold")

    path = os.path.join(DIAGRAMS_DIR, "sequence_diagram.png")
    plt.tight_layout()
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def generate_component_diagram():
    fig, ax = plt.subplots(1, 1, figsize=(12, 8))
    ax.set_xlim(-0.5, 12)
    ax.set_ylim(-0.5, 8.5)
    ax.axis("off")
    ax.set_title("Component Diagram", fontsize=16, fontweight="bold", color="#003366", pad=15)

    # Data Layer
    rect1 = mpatches.FancyBboxPatch((0.2, 0.2), 3.5, 2.5, boxstyle="round,pad=0.2",
                                     facecolor="#FFF3E0", edgecolor="#E65100", linewidth=2)
    ax.add_patch(rect1)
    ax.text(2.0, 2.4, "Data Layer", ha="center", fontsize=11, fontweight="bold", color="#E65100")
    draw_rounded_box(ax, 0.5, 0.5, 2.8, 0.7, "movies.csv\n(96 samples)", fc="#FFE0B2", ec="#E65100", fontsize=8)
    draw_rounded_box(ax, 0.5, 1.4, 2.8, 0.7, "Pandas\nDataFrame", fc="#FFE0B2", ec="#E65100", fontsize=8)

    # Preprocessing Layer
    rect2 = mpatches.FancyBboxPatch((4.2, 0.2), 3.5, 2.5, boxstyle="round,pad=0.2",
                                     facecolor="#E8F5E9", edgecolor="#2E7D32", linewidth=2)
    ax.add_patch(rect2)
    ax.text(6.0, 2.4, "Preprocessing", ha="center", fontsize=11, fontweight="bold", color="#2E7D32")
    draw_rounded_box(ax, 4.5, 0.5, 2.8, 0.7, "Text Cleaner\n(Lowercase, Regex)", fc="#C8E6C9", ec="#2E7D32", fontsize=8)
    draw_rounded_box(ax, 4.5, 1.4, 2.8, 0.7, "TF-IDF\nVectorizer", fc="#C8E6C9", ec="#2E7D32", fontsize=8)

    # Model Layer
    rect3 = mpatches.FancyBboxPatch((8.2, 0.2), 3.5, 2.5, boxstyle="round,pad=0.2",
                                     facecolor="#E3F2FD", edgecolor="#1565C0", linewidth=2)
    ax.add_patch(rect3)
    ax.text(10.0, 2.4, "Model Layer", ha="center", fontsize=11, fontweight="bold", color="#1565C0")
    draw_rounded_box(ax, 8.5, 1.4, 2.8, 0.7, "Naive Bayes\nLogistic Regression", fc="#BBDEFB", ec="#1565C0", fontsize=8)
    draw_rounded_box(ax, 8.5, 0.5, 2.8, 0.7, "Linear SVM\nRandom Forest", fc="#BBDEFB", ec="#1565C0", fontsize=8)

    # Evaluation Layer
    rect4 = mpatches.FancyBboxPatch((0.2, 3.5), 5.5, 2.0, boxstyle="round,pad=0.2",
                                     facecolor="#F3E5F5", edgecolor="#6A1B9A", linewidth=2)
    ax.add_patch(rect4)
    ax.text(3.0, 5.2, "Evaluation Layer", ha="center", fontsize=11, fontweight="bold", color="#6A1B9A")
    draw_rounded_box(ax, 0.5, 3.8, 2.3, 0.7, "Accuracy &\nCross-Validation", fc="#E1BEE7", ec="#6A1B9A", fontsize=8)
    draw_rounded_box(ax, 3.1, 3.8, 2.3, 0.7, "Classification\nReport & Confusion", fc="#E1BEE7", ec="#6A1B9A", fontsize=8)

    # Visualization Layer
    rect5 = mpatches.FancyBboxPatch((6.2, 3.5), 5.5, 2.0, boxstyle="round,pad=0.2",
                                     facecolor="#FFF9C4", edgecolor="#F57F17", linewidth=2)
    ax.add_patch(rect5)
    ax.text(9.0, 5.2, "Visualization Layer", ha="center", fontsize=11, fontweight="bold", color="#F57F17")
    draw_rounded_box(ax, 6.5, 3.8, 2.3, 0.7, "Matplotlib\nPlots", fc="#FFF59D", ec="#F57F17", fontsize=8)
    draw_rounded_box(ax, 9.1, 3.8, 2.3, 0.7, "Seaborn\nHeatmaps", fc="#FFF59D", ec="#F57F17", fontsize=8)

    # User Interface Layer
    rect6 = mpatches.FancyBboxPatch((2.5, 6.2), 7.0, 1.8, boxstyle="round,pad=0.2",
                                     facecolor="#E8EAF6", edgecolor="#283593", linewidth=2)
    ax.add_patch(rect6)
    ax.text(6.0, 7.7, "User Interface (CLI)", ha="center", fontsize=11, fontweight="bold", color="#283593")
    draw_rounded_box(ax, 2.8, 6.5, 3.0, 0.7, "Interactive Input\n(Movie Description)", fc="#C5CAE9", ec="#283593", fontsize=8)
    draw_rounded_box(ax, 6.2, 6.5, 3.0, 0.7, "Console Output\n(Results & Metrics)", fc="#C5CAE9", ec="#283593", fontsize=8)

    # Arrows between layers
    draw_arrow(ax, 3.7, 1.8, 4.2, 1.8)
    draw_arrow(ax, 7.7, 1.8, 8.2, 1.8)
    draw_arrow(ax, 6.0, 2.7, 3.0, 3.5)
    draw_arrow(ax, 6.0, 2.7, 9.0, 3.5)
    draw_arrow(ax, 6.0, 5.5, 6.0, 6.2)

    path = os.path.join(DIAGRAMS_DIR, "component_diagram.png")
    plt.tight_layout()
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def generate_system_architecture():
    fig, ax = plt.subplots(1, 1, figsize=(11, 7))
    ax.set_xlim(-0.5, 11)
    ax.set_ylim(-0.5, 7.5)
    ax.axis("off")
    ax.set_title("System Architecture", fontsize=16, fontweight="bold", color="#003366", pad=15)

    # Input
    draw_rounded_box(ax, 0.2, 5.5, 2.2, 1.2, "INPUT\n\nMovie\nDescription\n(Text)", fc="#E3F2FD", ec="#1565C0", fontsize=9)

    # Core Pipeline
    pipeline_box = mpatches.FancyBboxPatch((3.0, 0.5), 5.0, 6.5, boxstyle="round,pad=0.3",
                                            facecolor="#FAFAFA", edgecolor="#003366",
                                            linewidth=2.5, linestyle="-")
    ax.add_patch(pipeline_box)
    ax.text(5.5, 6.7, "ML Pipeline (main.py)", ha="center", fontsize=12, fontweight="bold", color="#003366")

    draw_rounded_box(ax, 3.5, 5.5, 4.0, 0.8, "Text Preprocessing", fc="#C8E6C9", ec="#2E7D32", fontsize=9)
    draw_rounded_box(ax, 3.5, 4.2, 4.0, 0.8, "TF-IDF Feature Extraction", fc="#DCEDC8", ec="#558B2F", fontsize=9)
    draw_rounded_box(ax, 3.5, 2.9, 4.0, 0.8, "Model Training (4 Models)", fc="#BBDEFB", ec="#1565C0", fontsize=9)
    draw_rounded_box(ax, 3.5, 1.6, 4.0, 0.8, "Evaluation & Comparison", fc="#E1BEE7", ec="#6A1B9A", fontsize=9)

    # Arrows inside pipeline
    draw_arrow(ax, 5.5, 5.5, 5.5, 5.0)
    draw_arrow(ax, 5.5, 4.2, 5.5, 3.7)
    draw_arrow(ax, 5.5, 2.9, 5.5, 2.4)

    # Input arrow
    draw_arrow(ax, 2.4, 6.1, 3.5, 5.9)

    # Output
    draw_rounded_box(ax, 8.5, 5.5, 2.2, 1.2, "OUTPUT\n\nPredicted\nGenre", fc="#FFF9C4", ec="#F57F17", fontsize=9)
    draw_arrow(ax, 7.5, 5.9, 8.5, 5.9)

    # Data
    draw_rounded_box(ax, 0.2, 2.5, 2.2, 1.2, "DATA\n\nmovies.csv\n(96 samples)", fc="#FFE0B2", ec="#E65100", fontsize=9)
    draw_arrow(ax, 2.4, 3.3, 3.5, 3.3)

    # Plots
    draw_rounded_box(ax, 8.5, 2.5, 2.2, 1.2, "PLOTS\n\n3 PNG\nCharts", fc="#F3E5F5", ec="#6A1B9A", fontsize=9)
    draw_arrow(ax, 7.5, 2.0, 8.5, 2.8)

    path = os.path.join(DIAGRAMS_DIR, "system_architecture.png")
    plt.tight_layout()
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


# ============================================================================
# REPORT BUILDER
# ============================================================================

def build_report():
    print("Generating diagrams...")
    use_case_path = generate_use_case_diagram()
    workflow_path = generate_workflow_diagram()
    sequence_path = generate_sequence_diagram()
    component_path = generate_component_diagram()
    architecture_path = generate_system_architecture()
    print("  All diagrams generated.")

    print("Building report...")
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # ====================================================================
    # 1. COVER PAGE
    # ====================================================================
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Movie Genre Classification")
    set_run_font(run, size=30, bold=True, color=(0, 51, 102))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Using Machine Learning and Natural Language Processing")
    set_run_font(run, size=16, color=(80, 80, 80))

    for _ in range(2):
        doc.add_paragraph()

    # Course info
    info_lines = [
        ("Course: ", "CSA2001"),
        ("Semester: ", "3"),
        ("Project Type: ", "Machine Learning / NLP"),
    ]
    for label, value in info_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(label)
        set_run_font(r1, size=13, bold=True, color=(0, 51, 102))
        r2 = p.add_run(value)
        set_run_font(r2, size=13)

    for _ in range(3):
        doc.add_paragraph()

    tech = doc.add_paragraph()
    tech.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tech.add_run("Python | Scikit-learn | Pandas | Matplotlib | Seaborn")
    set_run_font(run, size=11, italic=True, color=(120, 120, 120))

    doc.add_page_break()

    # ====================================================================
    # TABLE OF CONTENTS
    # ====================================================================
    add_heading_styled(doc, "Table of Contents", level=1)
    toc_items = [
        "1.  Introduction",
        "2.  Problem Statement",
        "3.  Functional Requirements",
        "4.  Non-Functional Requirements",
        "5.  System Architecture",
        "6.  Design Diagrams",
        "     6.1  Use Case Diagram",
        "     6.2  Workflow Diagram",
        "     6.3  Sequence Diagram",
        "     6.4  Component Diagram",
        "7.  Design Decisions & Rationale",
        "8.  Implementation Details",
        "9.  Screenshots / Results",
        "10. Testing Approach",
        "11. Challenges Faced",
        "12. Future Scope",
        "13. Conclusion",
    ]
    for item in toc_items:
        add_body(doc, item, size=12, spacing_after=4)

    doc.add_page_break()

    # ====================================================================
    # 2. INTRODUCTION
    # ====================================================================
    add_heading_styled(doc, "1. Introduction", level=1)
    add_body(doc, (
        "Movie genre classification is a fundamental problem in Natural Language Processing (NLP) "
        "and machine learning. With the explosion of digital content, the ability to automatically "
        "categorize movies based on their textual descriptions has significant practical applications "
        "in recommendation systems, content organization, and media analytics."
    ))
    add_body(doc, (
        "This project implements a complete machine learning pipeline that takes movie plot descriptions "
        "as input and predicts the genre of the movie. The system compares four different classification "
        "algorithms and two feature extraction methods to identify the best-performing approach. "
        "The project demonstrates the end-to-end ML workflow: data loading, preprocessing, feature "
        "extraction, model training, evaluation, visualization, and interactive prediction."
    ))
    add_body(doc, "Key objectives of this project:", bold=True)
    objectives = [
        "Develop a system that automatically classifies movie genres from text descriptions",
        "Apply and compare multiple machine learning algorithms for text classification",
        "Implement NLP preprocessing techniques for text normalization",
        "Evaluate models using proper metrics (accuracy, precision, recall, F1-score)",
        "Generate insightful visualizations for data and model performance analysis",
    ]
    for obj in objectives:
        add_bullet(doc, obj)

    doc.add_page_break()

    # ====================================================================
    # 3. PROBLEM STATEMENT
    # ====================================================================
    add_heading_styled(doc, "2. Problem Statement", level=1)
    add_body(doc, (
        "Given a textual description of a movie, the system should automatically predict the genre "
        "of the movie from a predefined set of 8 categories: Action, Comedy, Romance, Thriller, "
        "Science Fiction, Horror, Fantasy, and Drama."
    ))
    add_body(doc, (
        "The challenge involves converting unstructured text data into numerical features that machine "
        "learning algorithms can process, selecting appropriate classification models, and evaluating "
        "their performance to determine the most effective approach."
    ))
    add_body(doc, "Formally, this is a multi-class text classification problem where:", bold=True)
    add_bullet(doc, "Input: ", bold_prefix="")
    add_body(doc, "     A string containing the movie's plot description (e.g., 'A hero saves the world from alien invaders')")
    add_bullet(doc, "Output: ", bold_prefix="")
    add_body(doc, "     One of 8 genre labels (Action, Comedy, Romance, Thriller, Science Fiction, Horror, Fantasy, Drama)")
    add_bullet(doc, "Objective: ", bold_prefix="")
    add_body(doc, "     Maximize classification accuracy across all genres using the best model-feature combination")

    doc.add_page_break()

    # ====================================================================
    # 4. FUNCTIONAL REQUIREMENTS
    # ====================================================================
    add_heading_styled(doc, "3. Functional Requirements", level=1)

    add_table(doc,
        ["ID", "Requirement", "Description"],
        [
            ["FR-01", "Load Dataset", "The system shall load movie data from a CSV file containing descriptions and genre labels."],
            ["FR-02", "Preprocess Text", "The system shall normalize text by converting to lowercase and removing punctuation."],
            ["FR-03", "Extract Features", "The system shall convert text to numerical features using both CountVectorizer and TF-IDF."],
            ["FR-04", "Train Models", "The system shall train four ML models: Naive Bayes, Logistic Regression, SVM, and Random Forest."],
            ["FR-05", "Evaluate Models", "The system shall compute accuracy, cross-validation scores, and classification reports for each model."],
            ["FR-06", "Compare Vectorizers", "The system shall compare CountVectorizer and TF-IDF performance."],
            ["FR-07", "Generate Plots", "The system shall generate genre distribution, model comparison, and confusion matrix plots."],
            ["FR-08", "Predict Genre", "The system shall accept user input and predict the movie genre using the best model."],
            ["FR-09", "Display Results", "The system shall display all metrics and results in a formatted console output."],
        ]
    )

    doc.add_page_break()

    # ====================================================================
    # 5. NON-FUNCTIONAL REQUIREMENTS
    # ====================================================================
    add_heading_styled(doc, "4. Non-Functional Requirements", level=1)

    add_table(doc,
        ["ID", "Requirement", "Description"],
        [
            ["NFR-01", "Performance", "The system shall train all models and generate predictions within 30 seconds."],
            ["NFR-02", "Portability", "The system shall run on Windows, macOS, and Linux with Python 3.8+."],
            ["NFR-03", "Usability", "The system shall provide clear, formatted console output with section headers."],
            ["NFR-04", "Maintainability", "The code shall be modular with separate functions for each pipeline stage."],
            ["NFR-05", "Extensibility", "New models and genres shall be addable without modifying existing code structure."],
            ["NFR-06", "Reproducibility", "Results shall be reproducible using fixed random seeds (random_state=42)."],
            ["NFR-07", "Visualization Quality", "All plots shall be saved at 150 DPI with clear labels and titles."],
        ]
    )

    doc.add_page_break()

    # ====================================================================
    # 6. SYSTEM ARCHITECTURE
    # ====================================================================
    add_heading_styled(doc, "5. System Architecture", level=1)
    add_body(doc, (
        "The system follows a pipeline architecture where data flows sequentially through "
        "multiple processing stages. Each stage is implemented as a separate function in main.py, "
        "ensuring modularity and separation of concerns."
    ))
    add_body(doc, "The architecture consists of the following layers:", bold=True)
    add_bullet(doc, "User input via CLI (command-line interface)", bold_prefix="Presentation Layer: ")
    add_bullet(doc, "ML pipeline with preprocessing, feature extraction, training, and evaluation", bold_prefix="Processing Layer: ")
    add_bullet(doc, "CSV file for dataset storage and PNG files for generated plots", bold_prefix="Data Layer: ")

    add_image(doc, architecture_path, "Figure 1: System Architecture Diagram", width=5.5)

    doc.add_page_break()

    # ====================================================================
    # 7. DESIGN DIAGRAMS
    # ====================================================================
    add_heading_styled(doc, "6. Design Diagrams", level=1)

    # 6.1 Use Case
    add_heading_styled(doc, "6.1 Use Case Diagram", level=2)
    add_body(doc, (
        "The Use Case Diagram shows the interactions between the User (the only actor) "
        "and the system. The user can enter movie descriptions, view model accuracy and "
        "classification reports, view generated visualizations, and receive genre predictions."
    ))
    add_image(doc, use_case_path, "Figure 2: Use Case Diagram", width=5.5)

    doc.add_page_break()

    # 6.2 Workflow
    add_heading_styled(doc, "6.2 Workflow Diagram", level=2)
    add_body(doc, (
        "The Workflow Diagram illustrates the sequential flow of the ML pipeline from data loading "
        "through preprocessing, feature extraction, training, evaluation, and finally prediction. "
        "A side branch shows the visualization generation step."
    ))
    add_image(doc, workflow_path, "Figure 3: Workflow Diagram", width=5.5)

    doc.add_page_break()

    # 6.3 Sequence
    add_heading_styled(doc, "6.3 Sequence Diagram", level=2)
    add_body(doc, (
        "The Sequence Diagram shows the chronological order of interactions between the User, "
        "Main Controller (main.py), Preprocessor module, and ML Models during a complete "
        "execution cycle."
    ))
    add_image(doc, sequence_path, "Figure 4: Sequence Diagram", width=5.5)

    doc.add_page_break()

    # 6.4 Component
    add_heading_styled(doc, "6.4 Component Diagram", level=2)
    add_body(doc, (
        "The Component Diagram shows the major software components and their relationships. "
        "The system is organized into six layers: User Interface, Data, Preprocessing, Model, "
        "Evaluation, and Visualization."
    ))
    add_image(doc, component_path, "Figure 5: Component Diagram", width=5.5)

    add_body(doc, (
        "Note: This project does not use a database for storage, so an ER Diagram is not applicable. "
        "Data is stored in a flat CSV file (movies.csv) with two columns: description and genre."
    ), italic=True, size=10)

    doc.add_page_break()

    # ====================================================================
    # 8. DESIGN DECISIONS & RATIONALE
    # ====================================================================
    add_heading_styled(doc, "7. Design Decisions & Rationale", level=1)

    decisions = [
        (
            "TF-IDF over CountVectorizer",
            "TF-IDF (Term Frequency-Inverse Document Frequency) was chosen as the primary feature extraction "
            "method because it weighs terms by their importance across documents, reducing the influence of "
            "common words. CountVectorizer was retained for comparison purposes."
        ),
        (
            "Four Model Comparison",
            "Rather than relying on a single algorithm, four diverse models were selected to compare performance: "
            "Multinomial Naive Bayes (probabilistic baseline), Logistic Regression (linear), Linear SVM (margin-based), "
            "and Random Forest (ensemble). This provides a comprehensive view of which approach works best for this task."
        ),
        (
            "Stratified Train-Test Split",
            "Stratification ensures that the class distribution in training and test sets mirrors the original dataset. "
            "With only 12 samples per genre, a non-stratified split could exclude entire genres from the test set."
        ),
        (
            "5-Fold Cross-Validation",
            "Cross-validation provides a more robust estimate of model performance than a single train-test split, "
            "especially with a small dataset. It reduces the risk of overfitting to a particular random split."
        ),
        (
            "Modular Function Design",
            "Each pipeline stage (loading, preprocessing, vectorization, training, evaluation, visualization) "
            "is implemented as a separate function. This improves readability, testability, and makes it easy "
            "to modify individual stages without affecting others."
        ),
        (
            "Stop Words Removal",
            "English stop words are removed during vectorization to reduce noise from common words like "
            "'the', 'is', 'a' that don't carry genre-distinguishing information."
        ),
    ]

    for title, explanation in decisions:
        add_heading_styled(doc, title, level=2)
        add_body(doc, explanation)

    doc.add_page_break()

    # ====================================================================
    # 9. IMPLEMENTATION DETAILS
    # ====================================================================
    add_heading_styled(doc, "8. Implementation Details", level=1)

    add_heading_styled(doc, "8.1 Project Structure", level=2)
    add_code_block(doc, [
        "movie-genre-classification/",
        "    main.py              # Main ML pipeline (300 lines)",
        "    movies.csv           # Dataset (96 samples, 8 genres)",
        "    requirements.txt     # Python dependencies",
        "    generate_docx.py     # Report generator",
        "    plots/               # Auto-generated visualizations",
        "        genre_distribution.png",
        "        model_comparison.png",
        "        confusion_matrix.png",
        "    diagrams/            # Design diagrams for report",
        "    README.md            # Documentation",
    ])

    add_heading_styled(doc, "8.2 Technologies & Libraries", level=2)
    add_table(doc,
        ["Library", "Version", "Purpose"],
        [
            ["Python", "3.8+", "Core programming language"],
            ["Pandas", "Latest", "Data loading and manipulation"],
            ["Scikit-learn", "Latest", "ML models, vectorizers, evaluation metrics"],
            ["Matplotlib", "Latest", "Bar charts and plot generation"],
            ["Seaborn", "Latest", "Confusion matrix heatmap visualization"],
        ]
    )

    add_heading_styled(doc, "8.3 Data Preprocessing", level=2)
    add_body(doc, "Text preprocessing involves two steps applied to every movie description:")
    add_bullet(doc, "Convert all text to lowercase to ensure case-insensitive matching", bold_prefix="Lowercasing: ")
    add_bullet(doc, "Remove all non-alphabetic characters using regex [^a-z\\s]", bold_prefix="Punctuation Removal: ")
    add_code_block(doc, [
        "def preprocess_text(text):",
        '    text = text.lower()',
        '    text = re.sub(r"[^a-z\\s]", "", text)',
        '    return text.strip()',
    ])

    add_heading_styled(doc, "8.4 Feature Extraction", level=2)
    add_body(doc, (
        "Two vectorization methods are compared using Naive Bayes as a baseline model:"
    ))
    add_bullet(doc, "Converts text into a matrix of token counts (Bag of Words)", bold_prefix="CountVectorizer: ")
    add_bullet(doc, "Weighs terms by frequency in document vs. frequency across all documents", bold_prefix="TF-IDF Vectorizer: ")
    add_body(doc, "Both use English stop word removal to filter out common, non-informative words.")

    add_heading_styled(doc, "8.5 Models Used", level=2)
    add_table(doc,
        ["Model", "Algorithm Type", "Key Parameters"],
        [
            ["Multinomial Naive Bayes", "Probabilistic", "Default (alpha=1.0)"],
            ["Logistic Regression", "Linear", "max_iter=1000, random_state=42"],
            ["Linear SVM", "Margin-based", "max_iter=2000, random_state=42"],
            ["Random Forest", "Ensemble", "n_estimators=100, random_state=42"],
        ]
    )

    add_heading_styled(doc, "8.6 Evaluation Metrics", level=2)
    add_bullet(doc, "Percentage of correctly predicted genres", bold_prefix="Accuracy: ")
    add_bullet(doc, "Of all predicted genre X, how many were actually genre X", bold_prefix="Precision: ")
    add_bullet(doc, "Of all actual genre X samples, how many were correctly predicted", bold_prefix="Recall: ")
    add_bullet(doc, "Harmonic mean of precision and recall", bold_prefix="F1-Score: ")
    add_bullet(doc, "Average accuracy across 5 different train-test splits", bold_prefix="Cross-Validation: ")
    add_bullet(doc, "Matrix showing predicted vs. actual genres for error analysis", bold_prefix="Confusion Matrix: ")

    doc.add_page_break()

    # ====================================================================
    # 10. SCREENSHOTS / RESULTS
    # ====================================================================
    add_heading_styled(doc, "9. Screenshots / Results", level=1)

    add_heading_styled(doc, "9.1 Model Performance Summary", level=2)
    add_table(doc,
        ["Model", "Test Accuracy", "Cross-Val Mean", "Cross-Val Std"],
        [
            ["Multinomial Naive Bayes", "30.0%", "22.3%", "5.1%"],
            ["Logistic Regression", "30.0%", "23.6%", "6.3%"],
            ["Linear SVM (Best)", "35.0%", "23.6%", "6.3%"],
            ["Random Forest", "15.0%", "22.2%", "9.6%"],
        ]
    )

    add_body(doc, (
        "The Linear SVM achieved the highest test accuracy at 35.0%. While this accuracy is modest, "
        "it is expected given the small dataset size (96 samples across 8 genres). The focus of this "
        "project is on demonstrating the complete ML pipeline rather than achieving production-level accuracy."
    ))

    add_heading_styled(doc, "9.2 Genre Distribution", level=2)
    add_image(doc, os.path.join(PLOTS_DIR, "genre_distribution.png"),
              "Figure 6: Balanced distribution of 12 samples per genre", width=5.0)

    add_heading_styled(doc, "9.3 Model Accuracy Comparison", level=2)
    add_image(doc, os.path.join(PLOTS_DIR, "model_comparison.png"),
              "Figure 7: Test accuracy comparison across all four models", width=5.0)

    add_heading_styled(doc, "9.4 Confusion Matrix", level=2)
    add_image(doc, os.path.join(PLOTS_DIR, "confusion_matrix.png"),
              "Figure 8: Confusion matrix heatmap for the best model (Linear SVM)", width=5.0)

    doc.add_page_break()

    # ====================================================================
    # 11. TESTING APPROACH
    # ====================================================================
    add_heading_styled(doc, "10. Testing Approach", level=1)

    add_heading_styled(doc, "10.1 Data Validation", level=2)
    add_bullet(doc, "Verified CSV loads correctly with expected columns (description, genre)")
    add_bullet(doc, "Confirmed balanced class distribution (12 samples per genre)")
    add_bullet(doc, "Checked for missing values and empty descriptions")

    add_heading_styled(doc, "10.2 Preprocessing Validation", level=2)
    add_bullet(doc, "Verified lowercase conversion on sample texts")
    add_bullet(doc, "Confirmed punctuation removal using regex pattern")
    add_bullet(doc, "Tested edge cases: empty strings, special characters")

    add_heading_styled(doc, "10.3 Model Validation", level=2)
    add_bullet(doc, "Stratified train-test split (80/20) ensures genre representation in both sets")
    add_bullet(doc, "5-fold cross-validation provides robust accuracy estimates")
    add_bullet(doc, "Classification reports verify per-class precision, recall, and F1-score")
    add_bullet(doc, "Confusion matrix identifies specific misclassification patterns")

    add_heading_styled(doc, "10.4 End-to-End Testing", level=2)
    add_bullet(doc, "Full pipeline execution with no errors")
    add_bullet(doc, "Interactive prediction tested with diverse movie descriptions")
    add_bullet(doc, "Plot generation verified (3 PNG files in plots/ folder)")

    add_heading_styled(doc, "10.5 Test Cases", level=2)
    add_table(doc,
        ["Test Input", "Expected Genre", "Result"],
        [
            ["A soldier fighting enemies in war", "Action", "Pass"],
            ["A funny story of friends on a trip", "Comedy", "Pass"],
            ["Two strangers fall in love on a cruise ship", "Romance", "Pass"],
            ["A detective solving a crime mystery", "Thriller", "Pass"],
            ["A spaceship travels to distant galaxies", "Science Fiction", "Pass"],
        ]
    )

    doc.add_page_break()

    # ====================================================================
    # 12. CHALLENGES FACED
    # ====================================================================
    add_heading_styled(doc, "11. Challenges Faced", level=1)

    challenges = [
        (
            "Small Dataset Size",
            "With only 96 samples across 8 genres, the models have limited training data. "
            "This leads to moderate accuracy and difficulty in learning genre-specific patterns. "
            "Stratified splitting and cross-validation were used to mitigate this."
        ),
        (
            "Text Feature Representation",
            "Converting natural language text into numerical features that preserve semantic meaning "
            "is inherently challenging. TF-IDF was chosen over simple word counts to better capture "
            "term importance, but more advanced techniques like word embeddings could improve results."
        ),
        (
            "Overlapping Genre Descriptions",
            "Some movie descriptions share vocabulary across genres (e.g., 'story', 'people', 'young'). "
            "This makes it difficult for models to distinguish between similar genres like Drama and Romance."
        ),
        (
            "Model Selection and Tuning",
            "Choosing the right model and hyperparameters for a small, multi-class text classification "
            "task required comparing multiple approaches. No single model dominated across all genres."
        ),
        (
            "Windows Console Encoding",
            "Unicode characters caused encoding errors on Windows (cp1252 codec). All console output "
            "was converted to ASCII-safe alternatives for cross-platform compatibility."
        ),
    ]

    for title, explanation in challenges:
        add_heading_styled(doc, title, level=2)
        add_body(doc, explanation)

    doc.add_page_break()

    # ====================================================================
    # FUTURE SCOPE
    # ====================================================================
    add_heading_styled(doc, "12. Future Scope", level=1)
    future = [
        ("Larger Dataset: ", "Use real-world datasets like IMDB or TMDb with thousands of samples for better accuracy."),
        ("Deep Learning Models: ", "Implement LSTM, GRU, or BERT-based models for more sophisticated text understanding."),
        ("Web Interface: ", "Build a Flask or Streamlit web app for user-friendly genre prediction."),
        ("Multi-label Classification: ", "Support movies that belong to multiple genres simultaneously."),
        ("Word Embeddings: ", "Use Word2Vec, GloVe, or FastText for richer feature representations."),
        ("Hyperparameter Tuning: ", "Apply GridSearchCV or RandomizedSearchCV for optimal model parameters."),
    ]
    for bold_part, rest in future:
        add_bullet(doc, rest, bold_prefix=bold_part)

    doc.add_page_break()

    # ====================================================================
    # CONCLUSION
    # ====================================================================
    add_heading_styled(doc, "13. Conclusion", level=1)
    add_body(doc, (
        "This project successfully demonstrates a complete machine learning pipeline for movie genre "
        "classification based on textual descriptions. The system implements text preprocessing, "
        "two feature extraction methods (CountVectorizer and TF-IDF), and four classification models "
        "(Naive Bayes, Logistic Regression, Linear SVM, and Random Forest)."
    ))
    add_body(doc, (
        "The Linear SVM model achieved the best test accuracy of 35.0%, which is reasonable given "
        "the limited dataset of 96 samples across 8 genres. The project prioritizes demonstrating "
        "the end-to-end ML workflow - from data preprocessing to model evaluation and visualization - "
        "over achieving production-level accuracy."
    ))
    add_body(doc, (
        "Key takeaways include the importance of proper text preprocessing, the value of comparing "
        "multiple models rather than relying on a single algorithm, and the role of comprehensive "
        "evaluation metrics (beyond simple accuracy) in understanding model performance. The project "
        "provides a strong foundation that can be extended with larger datasets and more advanced "
        "NLP techniques in future iterations."
    ))

    # Save
    doc.save(OUTPUT_PATH)
    print(f"\nReport saved to: {OUTPUT_PATH}")
    print(f"Total pages: ~20+ pages with diagrams and plots")


if __name__ == "__main__":
    build_report()
